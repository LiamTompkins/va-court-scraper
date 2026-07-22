from __future__ import absolute_import
from __future__ import print_function
import os
import io
import re
import csv
import threading
import webbrowser
from datetime import datetime, timedelta

from flask import Flask, jsonify, request, Response
from sqlalchemy import create_engine, text
from rapidfuzz import fuzz, process

# Seconds without a heartbeat before an active task is considered stale.
# Matches the threshold used by task_watchdog.py.
STALE_THRESHOLD_SECONDS = 120
# Seconds without a presence heartbeat before a worker process is treated as
# gone and dropped from the dashboard.
WORKER_STALE_SECONDS = 90
PORT = 5000

app = Flask(__name__)
engine = create_engine("postgresql://" + os.environ['POSTGRES_DB'])


def ensure_schema():
    """Create the worker presence/target tables if they don't exist yet, so the
    dashboard and supervisor work even before a collector has built the schema."""
    ddl = [
        'CREATE TABLE IF NOT EXISTS workers ('
        ' worker_id VARCHAR PRIMARY KEY, court_type VARCHAR, status VARCHAR,'
        ' fips INTEGER, case_type VARCHAR, last_alive TIMESTAMP)',
        'CREATE TABLE IF NOT EXISTS worker_targets ('
        ' court_type VARCHAR PRIMARY KEY, desired_count INTEGER, updated_at TIMESTAMP)',
        'CREATE TABLE IF NOT EXISTS task_batches ('
        ' id SERIAL PRIMARY KEY, name VARCHAR, created_at TIMESTAMP, court_type VARCHAR,'
        ' case_type VARCHAR, startdate DATE, enddate DATE)',
        'CREATE TABLE IF NOT EXISTS retrieved_cases ('
        ' id BIGSERIAL PRIMARY KEY, batch_id INTEGER, court_type VARCHAR,'
        ' case_type VARCHAR, fips INTEGER, case_number VARCHAR, collected_at TIMESTAMP)',
    ]
    try:
        with engine.begin() as conn:
            for stmt in ddl:
                conn.execute(text(stmt))
    except Exception:
        pass
    # Add batch_id to any existing task tables (each in its own transaction so a
    # failure on a missing table doesn't abort the rest).
    for t in ['circuit_court_date_tasks', 'district_court_date_tasks',
              'circuit_court_active_date_tasks', 'district_court_active_date_tasks',
              'circuit_court_completed_date_tasks', 'district_court_completed_date_tasks']:
        try:
            with engine.begin() as conn:
                has_table = conn.execute(text(
                    "SELECT 1 FROM information_schema.tables WHERE table_name = :t"
                ), {'t': t}).scalar()
                if not has_table:
                    continue
                has_col = conn.execute(text(
                    "SELECT 1 FROM information_schema.columns "
                    "WHERE table_name = :t AND column_name = 'batch_id'"
                ), {'t': t}).scalar()
                if not has_col:
                    conn.execute(text('ALTER TABLE %s ADD COLUMN batch_id INTEGER' % t))
        except Exception:
            pass
    # Add columns that may predate their introduction on existing tables.
    for table, column, coltype in [
        ('task_batches', 'name', 'VARCHAR'),
    ]:
        try:
            with engine.begin() as conn:
                has_col = conn.execute(text(
                    "SELECT 1 FROM information_schema.columns "
                    "WHERE table_name = :t AND column_name = :c"
                ), {'t': table, 'c': column}).scalar()
                if not has_col:
                    conn.execute(text('ALTER TABLE %s ADD COLUMN %s %s' % (table, column, coltype)))
        except Exception:
            pass


ensure_schema()

# The court site becomes unstable past ~10 collectors, so cap the desired count.
MAX_WORKERS = 10
# Directory the supervisor writes per-worker logs to (relative to repo root).
LOG_DIR = 'worker_logs'

# Case tables are created by SQLAlchemy with mixed-case names. CASE_TABLE_NAMES
# holds the raw name (for pg_class lookups); queries that hit the table directly
# must double-quote it.
CASE_TABLE_NAMES = {
    'circuit': {'criminal': 'CircuitCriminalCase', 'civil': 'CircuitCivilCase'},
    'district': {'criminal': 'DistrictCriminalCase', 'civil': 'DistrictCivilCase'},
}

# Where to read each displayed field from, per court + case type. Column names
# differ between the four case tables, and civil parties live in their own
# tables, so the Collected Data view maps them here and joins live.
CASE_SOURCES = {
    ('district', 'criminal'): {
        'table': 'DistrictCriminalCase',
        'subtype': 'CaseType', 'judgement': 'FinalDisposition', 'date': 'FiledDate',
        'defendant': 'Name', 'defendant_attorney': 'DefenseAttorney',
        'parties': None,
    },
    ('circuit', 'criminal'): {
        'table': 'CircuitCriminalCase',
        'subtype': 'ChargeType', 'judgement': 'DispositionCode', 'date': 'Filed',
        'defendant': 'Defendant', 'defendant_attorney': 'DefendantsAttorney',
        'parties': None,
    },
    ('district', 'civil'): {
        'table': 'DistrictCivilCase',
        'subtype': 'CaseType', 'judgement': 'Judgment', 'date': 'FiledDate',
        'parties': ('DistrictCivilPlaintiff', 'DistrictCivilDefendant'),
    },
    ('circuit', 'civil'): {
        'table': 'CircuitCivilCase',
        'subtype': 'FilingType', 'judgement': 'Judgment', 'date': 'Filed',
        'parties': ('CircuitCivilPlaintiff', 'CircuitCivilDefendant'),
    },
}

# Cache of exact case counts, keyed by (court_type, case_type). Only used while a
# court has no active workers, where the count cannot change, so it stays valid
# until workers run again.
_exact_count_cache = {}


def scalar(conn, sql):
    """Run a COUNT-style query, returning 0 if the table doesn't exist yet."""
    try:
        return conn.execute(text(sql)).scalar() or 0
    except Exception:
        return 0


def case_count(conn, court_type, case_type, has_active_workers):
    """Return the case count for a table.

    While workers are active, use PostgreSQL's instant approximate row estimate
    (pg_class.reltuples) so the 1s refresh never scans the table. Once the court
    is idle, compute the exact COUNT(*) once and cache it (the count can't change
    with no workers running) so we still don't rescan every second.
    """
    table = CASE_TABLE_NAMES[court_type][case_type]
    key = (court_type, case_type)

    if has_active_workers:
        _exact_count_cache.pop(key, None)
        try:
            val = conn.execute(
                text('SELECT reltuples::bigint FROM pg_class WHERE relname = :n'),
                {'n': table}
            ).scalar()
            val = int(val) if val and val > 0 else 0
        except Exception:
            val = 0
        return {'count': val, 'approximate': True}

    if key not in _exact_count_cache:
        _exact_count_cache[key] = scalar(conn, 'SELECT COUNT(*) FROM "%s"' % table)
    return {'count': _exact_count_cache[key], 'approximate': False}


def collect_court_status(conn, court_type):
    active_table = '%s_court_active_date_tasks' % court_type
    pending_table = '%s_court_date_tasks' % court_type
    searched_table = '%s_court_dates_searched' % court_type

    active = []
    try:
        rows = conn.execute(text(
            'SELECT fips, casetype, startdate, enddate, last_alive '
            # Stable order: sorting by last_alive would reshuffle the rows on
            # every heartbeat.
            'FROM %s ORDER BY fips, casetype' % active_table
        ))
        now = datetime.now()
        for r in rows:
            last_alive = r[4]
            seconds_since = None
            stale = True
            if last_alive is not None:
                seconds_since = int((now - last_alive).total_seconds())
                stale = seconds_since > STALE_THRESHOLD_SECONDS
            active.append({
                'fips': str(r[0]).zfill(3),
                'case_type': r[1],
                'start_date': r[2].isoformat() if r[2] else None,
                'end_date': r[3].isoformat() if r[3] else None,
                'last_alive': last_alive.isoformat() if last_alive else None,
                'seconds_since': seconds_since,
                'stale': stale,
            })
    except Exception:
        pass

    completed_table = '%s_court_completed_date_tasks' % court_type
    completed = []
    try:
        rows = conn.execute(text(
            'SELECT fips, casetype, startdate, enddate, completed_at '
            'FROM %s ORDER BY completed_at DESC LIMIT 10' % completed_table
        ))
        for r in rows:
            completed.append({
                'fips': str(r[0]).zfill(3),
                'case_type': r[1],
                'start_date': r[2].isoformat() if r[2] else None,
                'end_date': r[3].isoformat() if r[3] else None,
                'completed_at': r[4].isoformat() if r[4] else None,
            })
    except Exception:
        pass

    # Idle workers come from the workers table (working ones are already shown
    # via their active task above).
    idle = []
    try:
        now = datetime.now()
        rows = conn.execute(text(
            "SELECT worker_id, last_alive FROM workers "
            "WHERE court_type = :ct AND status = 'idle' ORDER BY worker_id"
        ), {'ct': court_type})
        for r in rows:
            idle.append({
                'worker_id': r[0],
                'seconds_since': int((now - r[1]).total_seconds()) if r[1] else None,
            })
    except Exception:
        pass

    # Desired count (set via the dashboard) vs. actual running collectors.
    desired_count = 0
    running_count = 0
    try:
        d = conn.execute(text(
            'SELECT desired_count FROM worker_targets WHERE court_type = :ct'
        ), {'ct': court_type}).scalar()
        desired_count = int(d) if d is not None else 0
    except Exception:
        pass
    try:
        running_count = conn.execute(text(
            'SELECT COUNT(*) FROM workers WHERE court_type = :ct'
        ), {'ct': court_type}).scalar() or 0
    except Exception:
        pass

    has_active_workers = len(active) > 0
    return {
        'active': active,
        'idle': idle,
        'idle_count': len(idle),
        'desired_count': desired_count,
        'running_count': running_count,
        'completed': completed,
        'completed_count': scalar(conn, 'SELECT COUNT(*) FROM %s' % completed_table),
        'pending_count': scalar(conn, 'SELECT COUNT(*) FROM %s' % pending_table),
        'dates_searched': scalar(conn, 'SELECT COUNT(*) FROM %s' % searched_table),
        'cases': {
            'criminal': case_count(conn, court_type, 'criminal', has_active_workers),
            'civil': case_count(conn, court_type, 'civil', has_active_workers),
        },
    }


@app.route('/api/status')
def status():
    with engine.begin() as conn:
        # Drop workers that have stopped sending presence heartbeats.
        try:
            conn.execute(
                text('DELETE FROM workers WHERE last_alive < :cutoff'),
                {'cutoff': datetime.now() - timedelta(seconds=WORKER_STALE_SECONDS)}
            )
        except Exception:
            pass
        data = {
            'generated_at': datetime.now().isoformat(),
            'stale_threshold_seconds': STALE_THRESHOLD_SECONDS,
            'courts': {
                'district': collect_court_status(conn, 'district'),
                'circuit': collect_court_status(conn, 'circuit'),
            },
        }
    return jsonify(data)


COMPLETED_PER_PAGE = 10


@app.route('/api/completed')
def completed():
    try:
        page = max(0, int(request.args.get('page', 0)))
    except (TypeError, ValueError):
        page = 0

    rows = []
    with engine.connect() as conn:
        for court_type in ['district', 'circuit']:
            table = '%s_court_completed_date_tasks' % court_type
            try:
                result = conn.execute(text(
                    'SELECT fips, casetype, startdate, enddate, completed_at FROM %s' % table
                ))
                for r in result:
                    rows.append({
                        'court': court_type.capitalize(),
                        'fips': str(r[0]).zfill(3),
                        'case_type': r[1],
                        'start_date': r[2].isoformat() if r[2] else None,
                        'end_date': r[3].isoformat() if r[3] else None,
                        'completed_at': r[4].isoformat() if r[4] else None,
                        '_ts': r[4] or datetime.min,
                    })
            except Exception:
                pass

    rows.sort(key=lambda x: x['_ts'], reverse=True)
    total = len(rows)
    start = page * COMPLETED_PER_PAGE
    page_rows = rows[start:start + COMPLETED_PER_PAGE]
    for r in page_rows:
        del r['_ts']

    return jsonify({
        'page': page,
        'per_page': COMPLETED_PER_PAGE,
        'total': total,
        'tasks': page_rows,
    })


PENDING_PER_PAGE = 10


@app.route('/api/pending')
def pending():
    try:
        page = max(0, int(request.args.get('page', 0)))
    except (TypeError, ValueError):
        page = 0

    rows = []
    with engine.connect() as conn:
        for court_type in ['district', 'circuit']:
            table = '%s_court_date_tasks' % court_type
            try:
                result = conn.execute(text(
                    'SELECT id, fips, casetype, startdate, enddate FROM %s' % table
                ))
                for r in result:
                    rows.append({
                        'court': court_type.capitalize(),
                        'fips': str(r[1]).zfill(3),
                        'case_type': r[2],
                        'start_date': r[3].isoformat() if r[3] else None,
                        'end_date': r[4].isoformat() if r[4] else None,
                        '_id': r[0],
                    })
            except Exception:
                pass

    # No created-at column on tasks, so order by id (creation order), newest first.
    rows.sort(key=lambda x: x['_id'], reverse=True)
    total = len(rows)
    start = page * PENDING_PER_PAGE
    page_rows = rows[start:start + PENDING_PER_PAGE]
    for r in page_rows:
        del r['_id']

    return jsonify({
        'page': page,
        'per_page': PENDING_PER_PAGE,
        'total': total,
        'tasks': page_rows,
    })


@app.route('/api/tasks/create', methods=['POST'])
def create_tasks():
    data = request.get_json(force=True, silent=True) or {}
    court_type = (data.get('court_type') or '').strip()
    case_type = (data.get('case_type') or '').strip()
    name = (data.get('name') or '').strip() or None
    start_s = (data.get('start_date') or '').strip()
    end_s = (data.get('end_date') or '').strip()
    # fips may be a single value or a list of selected courts; empty means all.
    fips_raw = data.get('fips')
    if isinstance(fips_raw, list):
        selected = [str(f).strip() for f in fips_raw if str(f).strip()]
    elif fips_raw:
        selected = [str(fips_raw).strip()]
    else:
        selected = []

    if court_type not in ('circuit', 'district'):
        return jsonify({'ok': False, 'error': 'Invalid court level'}), 400
    if case_type not in ('criminal', 'civil'):
        return jsonify({'ok': False, 'error': 'Invalid case type'}), 400
    try:
        start_date = datetime.strptime(start_s, '%Y-%m-%d').date()
        end_date = datetime.strptime(end_s, '%Y-%m-%d').date()
    except ValueError:
        return jsonify({'ok': False, 'error': 'Invalid date (use YYYY-MM-DD)'}), 400
    # The form's "End (latest)" maps to start_date and "Start (earliest)" to
    # end_date. Collectors descend from start_date down to end_date, so the
    # most-recent date (start_date) must not precede the earliest (end_date).
    if start_date < end_date:
        return jsonify({'ok': False, 'error': 'Start date must be on or before end date'}), 400
    try:
        fips_list = [int(f) for f in selected]
    except ValueError:
        return jsonify({'ok': False, 'error': 'FIPS must be numeric'}), 400

    courts_table = '%s_courts' % court_type
    tasks_table = '%s_court_date_tasks' % court_type
    try:
        with engine.begin() as conn:
            if not fips_list:
                fips_list = [r[0] for r in conn.execute(text('SELECT fips FROM %s' % courts_table))]
            if not fips_list:
                return jsonify({'ok': False, 'error': 'No courts found - load courts first'}), 400
            batch_id = conn.execute(
                text('INSERT INTO task_batches (name, created_at, court_type, case_type, startdate, enddate) '
                     'VALUES (:nm, :now, :ct, :cs, :sd, :ed) RETURNING id'),
                {'nm': name, 'now': datetime.now(), 'ct': court_type, 'cs': case_type,
                 'sd': start_date, 'ed': end_date}
            ).scalar()
            for f in fips_list:
                conn.execute(
                    text('INSERT INTO %s (fips, startdate, enddate, casetype, batch_id) '
                         'VALUES (:fips, :sd, :ed, :ct, :bid)' % tasks_table),
                    {'fips': int(f), 'sd': start_date, 'ed': end_date, 'ct': case_type, 'bid': batch_id}
                )
    except Exception as e:
        return jsonify({'ok': False, 'error': str(e)}), 500
    return jsonify({'ok': True, 'created': len(fips_list), 'batch_id': batch_id})


@app.route('/api/workers/target', methods=['POST'])
def set_worker_target():
    data = request.get_json(force=True, silent=True) or {}
    court_type = (data.get('court_type') or '').strip()
    if court_type not in ('circuit', 'district'):
        return jsonify({'ok': False, 'error': 'Invalid court level'}), 400
    try:
        desired = int(data.get('desired_count'))
    except (TypeError, ValueError):
        return jsonify({'ok': False, 'error': 'desired_count must be a number'}), 400
    desired = max(0, min(MAX_WORKERS, desired))
    try:
        with engine.begin() as conn:
            conn.execute(text(
                'INSERT INTO worker_targets (court_type, desired_count, updated_at) '
                'VALUES (:ct, :dc, :now) '
                'ON CONFLICT (court_type) DO UPDATE SET '
                'desired_count = EXCLUDED.desired_count, updated_at = EXCLUDED.updated_at'
            ), {'ct': court_type, 'dc': desired, 'now': datetime.now()})
    except Exception as e:
        return jsonify({'ok': False, 'error': str(e)}), 500
    return jsonify({'ok': True, 'court_type': court_type, 'desired_count': desired})


def _safe_log_path(name):
    # Only allow plain *.log filenames inside LOG_DIR (no path traversal).
    if not name or '/' in name or '\\' in name or '..' in name or not name.endswith('.log'):
        return None
    path = os.path.join(LOG_DIR, name)
    return path if os.path.isfile(path) else None


@app.route('/api/logs')
def list_logs():
    items = []
    try:
        for n in os.listdir(LOG_DIR):
            if not n.endswith('.log'):
                continue
            try:
                st = os.stat(os.path.join(LOG_DIR, n))
                items.append({'name': n, 'size': st.st_size, 'modified': st.st_mtime})
            except Exception:
                pass
    except Exception:
        pass
    items.sort(key=lambda x: x['modified'], reverse=True)
    return jsonify({'logs': items})


@app.route('/api/logs/view')
def view_log():
    path = _safe_log_path(request.args.get('name', ''))
    if not path:
        return jsonify({'ok': False, 'error': 'Log not found'}), 404
    try:
        tail = max(1, min(2000, int(request.args.get('tail', 400))))
    except (TypeError, ValueError):
        tail = 400
    try:
        # Read only the tail end of the file so large logs stay cheap.
        with open(path, 'rb') as f:
            f.seek(0, os.SEEK_END)
            size = f.tell()
            f.seek(max(0, size - 256 * 1024))
            data = f.read().decode('utf-8', 'replace')
        content = '\n'.join(data.splitlines()[-tail:])
    except Exception as e:
        return jsonify({'ok': False, 'error': str(e)}), 500
    return jsonify({'ok': True, 'name': request.args.get('name', ''), 'content': content})


@app.route('/api/courts')
def list_courts():
    court_type = (request.args.get('court_type') or '').strip()
    if court_type not in ('circuit', 'district'):
        return jsonify({'courts': []})
    courts = []
    try:
        with engine.connect() as conn:
            rows = conn.execute(text('SELECT fips, name FROM %s_courts ORDER BY name' % court_type))
            for r in rows:
                courts.append({'fips': str(r[0]).zfill(3), 'name': r[1]})
    except Exception:
        pass
    return jsonify({'courts': courts})


BATCH_CASES_PER_PAGE = 50


@app.route('/api/batches')
def list_batches():
    batches = []
    try:
        with engine.connect() as conn:
            rows = conn.execute(text(
                'SELECT b.id, b.name, b.court_type, b.case_type, b.startdate, b.enddate, b.created_at, '
                '(SELECT COUNT(*) FROM retrieved_cases rc WHERE rc.batch_id = b.id) '
                'FROM task_batches b ORDER BY b.created_at DESC NULLS LAST, b.id DESC'
            ))
            for r in rows:
                batches.append({
                    'id': r[0],
                    'name': r[1],
                    'court_type': r[2],
                    'case_type': r[3],
                    'start_date': r[4].isoformat() if r[4] else None,
                    'end_date': r[5].isoformat() if r[5] else None,
                    'created_at': r[6].isoformat() if r[6] else None,
                    'cases': r[7] or 0,
                })
    except Exception:
        pass
    return jsonify({'batches': batches})


CASE_EXPORT_COLUMNS = [
    ('case_id', 'Case id'), ('case_number', 'Case number'), ('case_type', 'Category'),
    ('case_subtype', 'Case type'), ('judgement', 'Judgement'), ('fips', 'FIPS'),
    ('case_date', 'Case date'), ('plaintiff', 'Plaintiff'),
    ('plaintiff_attorney', 'Plaintiff attorney'), ('defendant', 'Defendant'),
    ('defendant_attorney', 'Defendant attorney'),
]


def fetch_batch_cases(conn, batch_id, sort=None, direction='ASC', limit=None, offset=0,
                      subtype_filter=None, judgement_filter=None):
    """Rows for a batch (and the total matching), joined live to the matching
    case table. Shared by the paginated view and the exports so they can't drift
    apart. limit=None returns every matching case. Optional case-type/judgement
    filters do a case-insensitive substring match on the joined columns.
    Returns (cases, total)."""
    batch = conn.execute(text(
        'SELECT court_type, case_type FROM task_batches WHERE id = :b'
    ), {'b': batch_id}).first()
    src = CASE_SOURCES.get((batch[0], batch[1])) if batch else None
    subtype_filter = (subtype_filter or '').strip()
    judgement_filter = (judgement_filter or '').strip()
    params = {'b': batch_id}

    if src is None:
        # No case table to join - the filters target joined columns, so a set
        # filter matches nothing.
        if subtype_filter or judgement_filter:
            return [], 0
        total = conn.execute(text(
            'SELECT COUNT(*) FROM retrieved_cases WHERE batch_id = :b'
        ), {'b': batch_id}).scalar() or 0
        tail = ''
        if limit is not None:
            params['lim'] = limit
            params['off'] = offset
            tail = ' LIMIT :lim OFFSET :off'
        rows = conn.execute(text(
            'SELECT case_number, fips, case_type FROM retrieved_cases '
            'WHERE batch_id = :b ORDER BY collected_at DESC NULLS LAST, id DESC' + tail
        ), params)
        return ([{'case_number': r[0], 'fips': str(r[1]).zfill(3), 'case_type': r[2],
                  'case_id': None, 'case_subtype': None, 'judgement': None,
                  'case_date': None, 'plaintiff': None, 'plaintiff_attorney': None,
                  'defendant': None, 'defendant_attorney': None} for r in rows], total)

    if src.get('parties'):
        # Civil: names/attorneys live in per-party tables.
        p_tbl, d_tbl = src['parties']
        def party(tbl, col):
            return ('(SELECT string_agg(x."%s", \'; \') FROM "%s" x WHERE x.case_id = c.id)'
                    % (col, tbl))
        plaintiff = party(p_tbl, 'Name')
        plaintiff_att = party(p_tbl, 'Attorney')
        defendant = party(d_tbl, 'Name')
        defendant_att = party(d_tbl, 'Attorney')
    else:
        # Criminal: single defendant on the case row, no plaintiff.
        plaintiff = 'NULL'
        plaintiff_att = 'NULL'
        defendant = 'c."%s"' % src['defendant']
        defendant_att = 'c."%s"' % src['defendant_attorney']

    join = ('FROM retrieved_cases rc '
            'LEFT JOIN "%s" c ON c.fips = rc.fips AND c."CaseNumber" = rc.case_number'
            % src['table'])
    where = 'rc.batch_id = :b'
    if subtype_filter:
        where += ' AND c."%s" ILIKE :subf' % src['subtype']
        params['subf'] = '%' + subtype_filter + '%'
    if judgement_filter:
        where += ' AND c."%s" ILIKE :judf' % src['judgement']
        params['judf'] = '%' + judgement_filter + '%'

    total = conn.execute(text(
        'SELECT COUNT(*) %s WHERE %s' % (join, where)), params).scalar() or 0

    # Map each sortable logical column to its real SQL expression. The logical
    # names (case_id, judgement, plaintiff, ...) are not actual columns in the
    # query, and some (fips) are ambiguous across the joined tables, so ordering
    # by the bare name raises a SQL error - which the endpoint would swallow and
    # render as "No cases recorded for this batch yet".
    sort_exprs = {
        'case_number': 'rc.case_number',
        'fips': 'rc.fips',
        'case_type': 'rc.case_type',
        'case_id': 'c.id',
        'case_subtype': 'c."%s"' % src['subtype'],
        'judgement': 'c."%s"' % src['judgement'],
        'case_date': 'c."%s"' % src['date'],
        'plaintiff': plaintiff,
        'plaintiff_attorney': plaintiff_att,
        'defendant': defendant,
        'defendant_attorney': defendant_att,
    }
    if sort in sort_exprs:
        order_by = '%s %s NULLS LAST, rc.id DESC' % (sort_exprs[sort], direction)
    else:
        order_by = 'rc.collected_at DESC NULLS LAST, rc.id DESC'

    tail = ''
    if limit is not None:
        params['lim'] = limit
        params['off'] = offset
        tail = ' LIMIT :lim OFFSET :off'

    sql = (
        'SELECT rc.case_number, rc.fips, rc.case_type, c.id, c."%s", c."%s", c."%s", '
        '%s, %s, %s, %s '
        '%s WHERE %s ORDER BY %s'
    ) % (src['subtype'], src['judgement'], src['date'],
         plaintiff, plaintiff_att, defendant, defendant_att, join, where, order_by) + tail

    cases = [{
        'case_number': r[0],
        'fips': str(r[1]).zfill(3),
        'case_type': r[2],
        'case_id': r[3],
        'case_subtype': r[4],
        'judgement': r[5],
        'case_date': r[6].isoformat() if r[6] else None,
        'plaintiff': r[7],
        'plaintiff_attorney': r[8],
        'defendant': r[9],
        'defendant_attorney': r[10],
    } for r in conn.execute(text(sql), params)]
    return cases, total


@app.route('/api/batches/<int:batch_id>/cases')
def batch_cases(batch_id):
    try:
        page = max(0, int(request.args.get('page', 0)))
    except (TypeError, ValueError):
        page = 0
    cases = []
    total = 0
    try:
        with engine.connect() as conn:
            direction = 'DESC' if (request.args.get('dir') or '').lower() == 'desc' else 'ASC'
            cases, total = fetch_batch_cases(
                conn, batch_id, request.args.get('sort') or '', direction,
                BATCH_CASES_PER_PAGE, page * BATCH_CASES_PER_PAGE,
                request.args.get('subtype'), request.args.get('judgement'))
    except Exception:
        pass
    return jsonify({'page': page, 'per_page': BATCH_CASES_PER_PAGE, 'total': total, 'cases': cases})


@app.route('/api/batches/<int:batch_id>/export')
def export_batch(batch_id):
    """Download a batch's cases as CSV (opens directly in Excel), honoring any
    sort/filter passed in the query string."""
    direction = 'DESC' if (request.args.get('dir') or '').lower() == 'desc' else 'ASC'
    try:
        with engine.connect() as conn:
            batch = conn.execute(text(
                'SELECT name FROM task_batches WHERE id = :b'
            ), {'b': batch_id}).first()
            rows, _ = fetch_batch_cases(
                conn, batch_id, request.args.get('sort') or '', direction,
                subtype_filter=request.args.get('subtype'),
                judgement_filter=request.args.get('judgement'))
    except Exception as e:
        return jsonify({'ok': False, 'error': str(e)}), 500

    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow([label for _key, label in CASE_EXPORT_COLUMNS])
    for row in rows:
        writer.writerow([row.get(key) for key, _label in CASE_EXPORT_COLUMNS])

    label = (batch[0] if batch and batch[0] else 'batch-%d' % batch_id)
    safe = re.sub(r'[^A-Za-z0-9._-]+', '-', label).strip('-') or ('batch-%d' % batch_id)
    return Response(
        buf.getvalue(),
        mimetype='text/csv',
        headers={'Content-Disposition': 'attachment; filename="%s-%d.csv"' % (safe, batch_id)}
    )


@app.route('/api/batches/<int:batch_id>/export.xlsx')
def export_batch_xlsx(batch_id):
    """Download every case in a batch as a real Excel workbook."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter
    except ImportError:
        return jsonify({'ok': False,
                        'error': 'openpyxl is not installed (pip install openpyxl)'}), 500

    direction = 'DESC' if (request.args.get('dir') or '').lower() == 'desc' else 'ASC'
    try:
        with engine.connect() as conn:
            batch = conn.execute(text(
                'SELECT name FROM task_batches WHERE id = :b'
            ), {'b': batch_id}).first()
            rows, _ = fetch_batch_cases(
                conn, batch_id, request.args.get('sort') or '', direction,
                subtype_filter=request.args.get('subtype'),
                judgement_filter=request.args.get('judgement'))
    except Exception as e:
        return jsonify({'ok': False, 'error': str(e)}), 500

    wb = Workbook()
    ws = wb.active
    ws.title = 'Cases'
    ws.append([label for _key, label in CASE_EXPORT_COLUMNS])
    for c in ws[1]:
        c.font = Font(bold=True)

    for row in rows:
        values = []
        for key, _label in CASE_EXPORT_COLUMNS:
            val = row.get(key)
            # Write real dates so Excel treats them as dates, not text. FIPS
            # stays a string so its leading zeros survive.
            if key == 'case_date' and val:
                try:
                    val = datetime.strptime(val, '%Y-%m-%d').date()
                except (TypeError, ValueError):
                    pass
            values.append(val)
        ws.append(values)

    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for i, (key, label) in enumerate(CASE_EXPORT_COLUMNS, start=1):
        longest = max([len(label)] + [len(str(r.get(key) or '')) for r in rows] or [0])
        ws.column_dimensions[get_column_letter(i)].width = min(max(longest + 2, 10), 40)

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    label = (batch[0] if batch and batch[0] else 'batch-%d' % batch_id)
    safe = re.sub(r'[^A-Za-z0-9._-]+', '-', label).strip('-') or ('batch-%d' % batch_id)
    return Response(
        bio.getvalue(),
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': 'attachment; filename="%s-%d.xlsx"' % (safe, batch_id)}
    )


# Conflict-check thresholds, mirrored from the elh-conflict-checker: a near
# match needs volunteer review, a near-perfect match is treated as the same
# person.
CONFLICT_REVIEW_THRESHOLD = 80
CONFLICT_EXACT_THRESHOLD = 99
# Party name fields (from fetch_batch_cases) to check, with display labels.
CONFLICT_PARTY_ROLES = [
    ('plaintiff', 'Plaintiff'),
    ('defendant', 'Defendant'),
    ('plaintiff_attorney', 'Plaintiff attorney'),
    ('defendant_attorney', 'Defendant attorney'),
]
_NAME_PUNCT = re.compile(r'[^a-z0-9 ]+')


def normalize_name(name):
    """Lowercase, strip punctuation, and collapse whitespace. Court records store
    names as "LAST, FIRST" while PracticePanther stores "First Last"; stripping
    the comma lets the token-based scorer match them regardless of word order."""
    if not name:
        return ''
    return ' '.join(_NAME_PUNCT.sub(' ', name.lower()).split())


def load_vplc_clients(conn):
    """VPLC's PracticePanther contacts, cached by the conflict checker in the
    shared 'pp_contact' table. Returns [] if the table doesn't exist yet (the
    conflict checker has never synced)."""
    try:
        rows = conn.execute(text(
            "SELECT display_name, first_name, last_name, role, is_adverse, elh_case_number "
            "FROM pp_contact WHERE display_name IS NOT NULL AND display_name <> ''"
        ))
    except Exception:
        return []
    clients = []
    for r in rows:
        clients.append({
            'display_name': r[0], 'first_name': r[1], 'last_name': r[2],
            'role': r[3], 'is_adverse': r[4], 'elh_case_number': r[5],
            'norm': normalize_name(r[0]),
        })
    return clients


def compute_batch_conflicts(conn, batch_id):
    """Scan every party in a batch's cases against VPLC's PracticePanther client
    list, flagging fuzzy name matches (potential conflicts of interest). Shared
    by the JSON endpoint and the Excel export so they can't drift apart.
    Returns (conflicts, clients_checked, cases_scanned, note); note is set (with
    conflicts empty) when there are no VPLC contacts to check against."""
    clients = load_vplc_clients(conn)
    if not clients:
        return [], 0, 0, ("No VPLC (PracticePanther) contacts found in the shared "
                          "database. Run the conflict checker's PracticePanther sync first.")
    cases, _ = fetch_batch_cases(conn, batch_id)  # limit=None: all cases

    # Collect every non-empty party name, keyed by its normalized form, so a name
    # shared across cases is scored once (rapidfuzz is C-fast, but this keeps a
    # big batch cheap).
    occurrences = {}
    for c in cases:
        for key, label in CONFLICT_PARTY_ROLES:
            raw = c.get(key)
            norm = normalize_name(raw)
            if not norm:
                continue
            occurrences.setdefault(norm, []).append({
                'case_number': c['case_number'], 'fips': c['fips'],
                'role': label, 'name': raw, 'case': c,
            })

    client_norms = [cl['norm'] for cl in clients]
    conflicts = []
    for norm, occ_list in occurrences.items():
        # token_sort_ratio is order-insensitive, so a normalized "smith john"
        # matches a client stored as "john smith".
        matches = process.extract(
            norm, client_norms, scorer=fuzz.token_sort_ratio,
            score_cutoff=CONFLICT_REVIEW_THRESHOLD, limit=None)
        if not matches:
            continue
        for occ in occ_list:
            for _matched, score, idx in matches:
                cl = clients[idx]
                conflicts.append({
                    'score': int(round(score)),
                    'level': 'exact' if score >= CONFLICT_EXACT_THRESHOLD else 'review',
                    'case_number': occ['case_number'], 'fips': occ['fips'],
                    'party_role': occ['role'], 'party_name': occ['name'],
                    'client_name': cl['display_name'], 'client_role': cl['role'],
                    'client_adverse': bool(cl['is_adverse']),
                    'elh_case_number': cl['elh_case_number'],
                    # Full Collected-Data case record, carried through to the
                    # exports so each flagged row keeps all of the case's fields.
                    'case': occ['case'],
                })

    conflicts.sort(key=lambda x: x['score'], reverse=True)
    return conflicts, len(clients), len(cases), None


# The conflict exports lead with the match details, then append the flagged
# case's full Collected-Data record (CASE_EXPORT_COLUMNS) so every field already
# shown in the batch's case table is preserved. case_number/fips are omitted here
# because they come from the case columns.
CONFLICT_MATCH_COLUMNS = [
    ('score', 'Match %'), ('level', 'Level'), ('party_role', 'Party role'),
    ('party_name', 'Party name'), ('client_name', 'VPLC client'),
    ('client_role', 'Client role'), ('client_adverse', 'Adverse'),
    ('elh_case_number', 'ELH case'),
]


def _conflict_cell(row, key):
    """Display value for one conflict-match field (booleans rendered as Yes/No)."""
    val = row.get(key)
    if key == 'client_adverse':
        return 'Yes' if val else 'No'
    return val


def conflict_export_header():
    return ([label for _k, label in CONFLICT_MATCH_COLUMNS] +
            [label for _k, label in CASE_EXPORT_COLUMNS])


def conflict_export_row(cf):
    """One export row: match details followed by the flagged case's full record."""
    case = cf.get('case') or {}
    return ([_conflict_cell(cf, k) for k, _l in CONFLICT_MATCH_COLUMNS] +
            [case.get(k) for k, _l in CASE_EXPORT_COLUMNS])


@app.route('/api/batches/<int:batch_id>/conflicts')
def batch_conflicts(batch_id):
    try:
        with engine.connect() as conn:
            conflicts, clients_checked, cases_scanned, note = \
                compute_batch_conflicts(conn, batch_id)
    except Exception as e:
        return jsonify({'ok': False, 'error': str(e)}), 500
    resp = {'ok': True, 'clients_checked': clients_checked,
            'cases_scanned': cases_scanned, 'conflicts': conflicts}
    if note:
        resp['note'] = note
    return jsonify(resp)


@app.route('/api/batches/<int:batch_id>/conflicts.xlsx')
def export_conflicts_xlsx(batch_id):
    """Download a batch's potential conflicts as a real Excel workbook."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter
    except ImportError:
        return jsonify({'ok': False,
                        'error': 'openpyxl is not installed (pip install openpyxl)'}), 500

    try:
        with engine.connect() as conn:
            batch = conn.execute(text(
                'SELECT name FROM task_batches WHERE id = :b'), {'b': batch_id}).first()
            conflicts, _clients, _cases, _note = compute_batch_conflicts(conn, batch_id)
    except Exception as e:
        return jsonify({'ok': False, 'error': str(e)}), 500

    wb = Workbook()
    ws = wb.active
    ws.title = 'Conflicts'
    header = conflict_export_header()
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
    data_rows = [conflict_export_row(cf) for cf in conflicts]
    for row in data_rows:
        ws.append(row)

    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for i, label in enumerate(header, start=1):
        longest = max([len(label)] +
                      [len(str(r[i - 1] if r[i - 1] is not None else '')) for r in data_rows])
        ws.column_dimensions[get_column_letter(i)].width = min(max(longest + 2, 10), 40)

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    label = (batch[0] if batch and batch[0] else 'batch-%d' % batch_id)
    safe = re.sub(r'[^A-Za-z0-9._-]+', '-', label).strip('-') or ('batch-%d' % batch_id)
    return Response(
        bio.getvalue(),
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': 'attachment; filename="%s-conflicts-%d.xlsx"' % (safe, batch_id)}
    )


@app.route('/api/batches/<int:batch_id>/conflicts.csv')
def export_conflicts_csv(batch_id):
    """Download a batch's potential conflicts as CSV (opens directly in Excel)."""
    try:
        with engine.connect() as conn:
            batch = conn.execute(text(
                'SELECT name FROM task_batches WHERE id = :b'), {'b': batch_id}).first()
            conflicts, _clients, _cases, _note = compute_batch_conflicts(conn, batch_id)
    except Exception as e:
        return jsonify({'ok': False, 'error': str(e)}), 500

    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(conflict_export_header())
    for cf in conflicts:
        writer.writerow(conflict_export_row(cf))

    label = (batch[0] if batch and batch[0] else 'batch-%d' % batch_id)
    safe = re.sub(r'[^A-Za-z0-9._-]+', '-', label).strip('-') or ('batch-%d' % batch_id)
    return Response(
        buf.getvalue(),
        mimetype='text/csv',
        headers={'Content-Disposition': 'attachment; filename="%s-conflicts-%d.csv"' % (safe, batch_id)}
    )


@app.route('/')
def index():
    return PAGE


@app.route('/data')
def data_page():
    return DATA_PAGE


PAGE = """<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<title>VA Court Scraper - Worker Dashboard</title>
<style>
  body { font-family: -apple-system, Segoe UI, Roboto, sans-serif; margin: 24px; background: #0f1115; color: #e6e6e6; }
  h1 { font-size: 20px; margin: 0 0 4px; }
  .top { display: flex; align-items: center; justify-content: space-between; }
  .logs-btn { background: #21262d; color: #e6e6e6; border: 1px solid #30363d; border-radius: 6px; padding: 6px 12px; cursor: pointer; font-size: 13px; }
  .logs-btn:hover { background: #30363d; }
  .modal { position: fixed; inset: 0; background: rgba(0,0,0,.6); display: flex; align-items: center; justify-content: center; z-index: 50; }
  .modal-box { background: #171a21; border: 1px solid #262b36; border-radius: 8px; width: 80%; max-width: 1000px; height: 80%; display: flex; flex-direction: column; }
  .modal-head { display: flex; align-items: center; justify-content: space-between; padding: 12px 16px; border-bottom: 1px solid #262b36; }
  .modal-head h2 { margin: 0; font-size: 16px; }
  .modal-head button { background: none; border: none; color: #8a8f98; font-size: 22px; cursor: pointer; line-height: 1; }
  .modal-body { display: flex; flex: 1; min-height: 0; }
  .log-list { width: 280px; border-right: 1px solid #262b36; overflow-y: auto; }
  .log-item { padding: 8px 12px; cursor: pointer; border-bottom: 1px solid #1f242c; }
  .log-item:hover { background: #1f242c; }
  .log-item.active { background: #21262d; }
  .log-item .ln { font-size: 12px; color: #e6e6e6; word-break: break-all; }
  .log-item .lm { font-size: 11px; color: #8a8f98; }
  .log-view { flex: 1; margin: 0; padding: 12px; overflow: auto; font-family: ui-monospace, Menlo, Consolas, monospace; font-size: 12px; white-space: pre-wrap; color: #cdd3dc; }
  .meta { color: #8a8f98; font-size: 13px; margin-bottom: 20px; }
  .grid { display: flex; gap: 24px; flex-wrap: wrap; }
  .court { flex: 1; min-width: 420px; background: #171a21; border: 1px solid #262b36; border-radius: 8px; padding: 16px; }
  .completed-box { flex: 0 1 calc(50% - 12px); max-width: calc(50% - 12px); }
  .scheduler { background: #171a21; border: 1px solid #262b36; border-radius: 8px; padding: 16px; }
  .scheduler h2 { font-size: 16px; margin: 0 0 12px; }
  .scheduler .row { display: flex; gap: 12px; align-items: flex-end; flex-wrap: wrap; }
  .scheduler label { display: flex; flex-direction: column; font-size: 11px; color: #8a8f98; text-transform: uppercase; letter-spacing: .04em; gap: 4px; }
  .scheduler select, .scheduler input { background: #0f1115; color: #e6e6e6; border: 1px solid #30363d; border-radius: 6px; padding: 6px 8px; font-size: 13px; }
  .scheduler button { background: #238636; color: #fff; border: 1px solid #2ea043; border-radius: 6px; padding: 7px 16px; cursor: pointer; font-size: 13px; }
  .scheduler button:hover { background: #2ea043; }
  .scheduler .msg { margin-left: 8px; font-size: 13px; color: #8a8f98; }
  .dropdown { position: relative; }
  .scheduler .dd-toggle { background: #0f1115; color: #e6e6e6; border: 1px solid #30363d; border-radius: 6px; padding: 6px 8px; font-size: 13px; cursor: pointer; text-align: left; min-width: 130px; }
  .scheduler .dd-toggle:hover { background: #0f1115; }
  .scheduler .dd-toggle::after { content: ' ▾'; color: #8a8f98; }
  .dd-panel { position: absolute; top: 100%; left: 0; margin-top: 4px; display: none; z-index: 20; box-shadow: 0 4px 16px rgba(0,0,0,.4); }
  .dd-panel.open { display: block; }
  .court-checks { max-height: 200px; overflow-y: auto; border: 1px solid #30363d; border-radius: 6px; padding: 6px 8px; background: #171a21; min-width: 240px; }
  .court-checks label { display: flex; flex-direction: row; align-items: center; gap: 6px; font-size: 13px; text-transform: none; letter-spacing: normal; color: #e6e6e6; padding: 1px 0; white-space: nowrap; }
  .court-checks input { margin: 0; }
  .court-checks .empty { font-size: 13px; }
  .court h2 { font-size: 16px; margin: 0 0 12px; text-transform: capitalize; }
  .court-head { display: flex; align-items: center; justify-content: space-between; }
  .court-head h2 { margin: 0 0 12px; }
  .worker-ctl { display: flex; align-items: center; gap: 6px; font-size: 13px; color: #8a8f98; }
  .worker-ctl .run { color: #e6e6e6; }
  .worker-ctl .lbl { font-size: 11px; text-transform: uppercase; letter-spacing: .04em; }
  .worker-ctl button { background: #21262d; color: #e6e6e6; border: 1px solid #30363d; border-radius: 6px; width: 24px; height: 24px; cursor: pointer; font-size: 14px; line-height: 1; }
  .worker-ctl button:hover { background: #30363d; }
  .court h3.section { font-size: 12px; color: #8a8f98; text-transform: uppercase; letter-spacing: .04em; margin: 18px 0 8px; }
  .stats { display: flex; gap: 16px; margin-bottom: 14px; flex-wrap: wrap; }
  .stat { background: #0f1115; border-radius: 6px; padding: 8px 12px; min-width: 90px; }
  .stat .n { font-size: 20px; font-weight: 600; }
  .stat .l { font-size: 11px; color: #8a8f98; text-transform: uppercase; letter-spacing: .04em; }
  table { width: 100%; border-collapse: collapse; font-size: 13px; }
  th, td { text-align: left; padding: 6px 8px; border-bottom: 1px solid #262b36; }
  th { color: #8a8f98; font-weight: 500; font-size: 11px; text-transform: uppercase; }
  .ok { color: #3fb950; }
  .stale { color: #f85149; font-weight: 600; }
  .empty { color: #8a8f98; font-style: italic; padding: 8px; }
  .dot { display: inline-block; width: 8px; height: 8px; border-radius: 50%; margin-right: 6px; }
  .dot.ok { background: #3fb950; }
  .dot.stale { background: #f85149; }
  .dot.idle { background: #8a8f98; }
  .muted { color: #8a8f98; }
  .pager { display: flex; align-items: center; gap: 12px; margin-top: 12px; font-size: 13px; color: #8a8f98; }
  .pager button { background: #21262d; color: #e6e6e6; border: 1px solid #30363d; border-radius: 6px; padding: 5px 12px; cursor: pointer; }
  .pager button:hover:not([disabled]) { background: #30363d; }
  .pager button[disabled] { opacity: .4; cursor: default; }
</style>
</head>
<body>
  <div class="top">
    <h1>Virginia Court Scraper &mdash; Worker Dashboard</h1>
    <span>
      <a class="logs-btn" href="/data" style="text-decoration:none;">Collected data</a>
      <button class="logs-btn" onclick="openLogs()">Worker logs</button>
    </span>
  </div>
  <div class="meta" id="meta">Loading&hellip;</div>
  <div class="modal" id="logs-modal" style="display:none;">
    <div class="modal-box">
      <div class="modal-head">
        <h2>Worker logs</h2>
        <button onclick="closeLogs()">&times;</button>
      </div>
      <div class="modal-body">
        <div class="log-list" id="log-list"></div>
        <pre class="log-view" id="log-view">Select a log to view&hellip;</pre>
      </div>
    </div>
  </div>
  <div class="grid" id="grid"></div>
  <div class="scheduler" style="margin-top:24px;">
    <h2>Task Scheduler</h2>
    <div class="row">
      <label>Name (optional)
        <input type="text" id="sch-name" placeholder="batch name" size="16">
      </label>
      <label>Court level
        <select id="sch-court" onchange="loadCourts()"><option value="district">district</option><option value="circuit">circuit</option></select>
      </label>
      <label>Case type
        <select id="sch-case"><option value="criminal">criminal</option><option value="civil">civil</option></select>
      </label>
      <label>Start (earliest)
        <input type="date" id="sch-end">
      </label>
      <label>End (latest)
        <input type="date" id="sch-start">
      </label>
      <label>Courts (none = all)
        <div class="dropdown">
          <button type="button" class="dd-toggle" id="sch-courts-toggle" onclick="toggleCourts(event)">All courts</button>
          <div class="dd-panel court-checks" id="sch-courts"></div>
        </div>
      </label>
      <button onclick="createTasks()">Create tasks</button>
      <span class="msg" id="sch-msg"></span>
    </div>
  </div>
  <div class="grid" id="grid-completed" style="margin-top:24px;">
    <div class="court" id="completed-card"></div>
    <div class="court" id="pending-card"></div>
  </div>
<script>
function fmtAgo(s) {
  if (s === null) return 'no heartbeat';
  if (s < 60) return s + 's ago';
  var m = Math.floor(s / 60);
  return m + 'm ' + (s % 60) + 's ago';
}
var currentLog = null;
var logsTimer = null;
function openLogs() {
  document.getElementById('logs-modal').style.display = 'flex';
  loadLogList();
}
function closeLogs() {
  document.getElementById('logs-modal').style.display = 'none';
  if (logsTimer) { clearTimeout(logsTimer); logsTimer = null; }
  currentLog = null;
}
function loadLogList() {
  fetch('/api/logs').then(function(r){ return r.json(); }).then(function(d){
    var html = (d.logs || []).map(function(l){
      var when = new Date(l.modified * 1000).toLocaleString();
      var kb = (l.size / 1024).toFixed(1);
      var active = (l.name === currentLog) ? ' active' : '';
      return '<div class="log-item' + active + '" onclick="viewLog(&#39;' + l.name + '&#39;)">' +
        '<div class="ln">' + l.name + '</div>' +
        '<div class="lm">' + kb + ' KB · ' + when + '</div></div>';
    }).join('');
    document.getElementById('log-list').innerHTML = html || '<div class="empty">No logs yet</div>';
  });
}
function viewLog(name) {
  currentLog = name;
  loadLogList();
  refreshLogView();
}
function refreshLogView() {
  if (!currentLog) return;
  fetch('/api/logs/view?name=' + encodeURIComponent(currentLog) + '&tail=400')
    .then(function(r){ return r.json(); }).then(function(d){
      if (!d.ok) return;
      var v = document.getElementById('log-view');
      var atBottom = v.scrollTop + v.clientHeight >= v.scrollHeight - 20;
      v.textContent = d.content || '(empty)';
      if (atBottom) v.scrollTop = v.scrollHeight;
    });
  if (logsTimer) clearTimeout(logsTimer);
  logsTimer = setTimeout(refreshLogView, 3000);
}
function setTarget(court, desired) {
  if (desired < 0) desired = 0;
  fetch('/api/workers/target', {
    method: 'POST',
    headers: {'Content-Type': 'application/json'},
    body: JSON.stringify({court_type: court, desired_count: desired})
  }).then(function(r){ return r.json(); }).then(function(d){ if (d.ok) refresh(); });
}
function toggleCourts(e) {
  e.stopPropagation();
  document.getElementById('sch-courts').classList.toggle('open');
}
function updateCourtsLabel() {
  var n = document.querySelectorAll('#sch-courts input:checked').length;
  document.getElementById('sch-courts-toggle').textContent = n === 0 ? 'All courts' : (n + ' selected');
}
function loadCourts() {
  var court = document.getElementById('sch-court').value;
  var box = document.getElementById('sch-courts');
  fetch('/api/courts?court_type=' + court).then(function(r){ return r.json(); }).then(function(d){
    var rows = (d.courts || []).map(function(c){
      return '<label><input type="checkbox" value="' + c.fips + '"> ' + c.name + ' (' + c.fips + ')</label>';
    }).join('');
    box.innerHTML = rows || '<span class="empty">No courts loaded</span>';
    Array.prototype.forEach.call(box.querySelectorAll('input[type=checkbox]'), function(cb){
      cb.addEventListener('change', updateCourtsLabel);
    });
    updateCourtsLabel();
  });
}
// Close the courts dropdown when clicking outside it.
document.addEventListener('click', function(e){
  var panel = document.getElementById('sch-courts');
  var toggle = document.getElementById('sch-courts-toggle');
  if (panel && !panel.contains(e.target) && e.target !== toggle) {
    panel.classList.remove('open');
  }
});
function createTasks() {
  var msg = document.getElementById('sch-msg');
  var checked = document.querySelectorAll('#sch-courts input:checked');
  var fips = Array.prototype.map.call(checked, function(c){ return c.value; });
  var body = {
    name: document.getElementById('sch-name').value,
    court_type: document.getElementById('sch-court').value,
    case_type: document.getElementById('sch-case').value,
    start_date: document.getElementById('sch-start').value,
    end_date: document.getElementById('sch-end').value,
    fips: fips
  };
  if (!body.start_date || !body.end_date) { msg.textContent = 'Pick start and end dates.'; return; }
  msg.textContent = 'Creating…';
  fetch('/api/tasks/create', {
    method: 'POST',
    headers: {'Content-Type': 'application/json'},
    body: JSON.stringify(body)
  }).then(function(r){ return r.json(); }).then(function(d){
    if (d.ok) { msg.textContent = 'Created ' + d.created + ' task(s).'; refresh(); }
    else { msg.textContent = 'Error: ' + d.error; }
  }).catch(function(e){ msg.textContent = 'Error: ' + e; });
}
function fmtCount(c) {
  // c is {count, approximate}; show ~ while the figure is an estimate
  return (c.approximate ? '~' : '') + c.count.toLocaleString();
}
function fmtWhen(iso) {
  if (!iso) return '';
  var d = new Date(iso);
  var secs = Math.floor((Date.now() - d.getTime()) / 1000);
  if (secs < 60) return secs + 's ago';
  if (secs < 3600) return Math.floor(secs / 60) + 'm ago';
  if (secs < 86400) {
    var h = Math.floor(secs / 3600);
    var m = Math.floor((secs % 3600) / 60);
    return h + 'h ' + m + 'm ago';
  }
  return Math.round(secs / 86400) + 'd ago';
}
function courtCard(name, c) {
  var rows = c.active.map(function(t) {
    var cls = t.stale ? 'stale' : 'ok';
    return '<tr>' +
      '<td><span class="dot ' + cls + '"></span>working</td>' +
      '<td>' + t.fips + '</td>' +
      '<td>' + t.case_type + '</td>' +
      '<td>' + (t.end_date || '') + ' → ' + (t.start_date || '') + '</td>' +
      '<td class="' + cls + '">' + fmtAgo(t.seconds_since) + '</td>' +
      '</tr>';
  }).join('');
  var idleRows = (c.idle || []).map(function(w) {
    return '<tr>' +
      '<td><span class="dot idle"></span>idle</td>' +
      '<td colspan="3" class="muted">' + w.worker_id + '</td>' +
      '<td>' + fmtAgo(w.seconds_since) + '</td>' +
      '</tr>';
  }).join('');
  var allRows = rows + idleRows;
  var table = allRows
    ? '<table><tr><th>Status</th><th>FIPS</th><th>Type</th><th>Date range</th><th>Last heartbeat</th></tr>' + allRows + '</table>'
    : '<div class="empty">No workers</div>';

  var ct = name.toLowerCase();
  var desired = c.desired_count || 0;
  var workerCtl = '<div class="worker-ctl">' +
    '<span class="run">' + (c.running_count || 0) + ' running</span> / ' +
    '<button onclick="setTarget(&#39;' + ct + '&#39;,' + (desired - 1) + ')">&minus;</button>' +
    '<b>' + desired + '</b>' +
    '<button onclick="setTarget(&#39;' + ct + '&#39;,' + (desired + 1) + ')">+</button>' +
    '<span class="lbl">desired</span></div>';

  return '<div class="court">' +
    '<div class="court-head"><h2>' + name + '</h2>' + workerCtl + '</div>' +
    '<div class="stats">' +
      '<div class="stat"><div class="n">' + c.pending_count + '</div><div class="l">Pending</div></div>' +
      '<div class="stat"><div class="n">' + c.completed_count + '</div><div class="l">Completed</div></div>' +
      '<div class="stat"><div class="n">' + c.dates_searched + '</div><div class="l">Dates done</div></div>' +
      '<div class="stat"><div class="n">' + fmtCount(c.cases.criminal) + '</div><div class="l">Criminal cases</div></div>' +
      '<div class="stat"><div class="n">' + fmtCount(c.cases.civil) + '</div><div class="l">Civil cases</div></div>' +
    '</div>' + table + '</div>';
}
var completedPage = 0;
var lastStatus = null;
function renderCompleted(data) {
  var doneRows = data.tasks.map(function(t) {
    return '<tr>' +
      '<td>' + t.court + '</td>' +
      '<td>' + t.fips + '</td>' +
      '<td>' + t.case_type + '</td>' +
      '<td>' + (t.end_date || '') + ' → ' + (t.start_date || '') + '</td>' +
      '<td>' + fmtWhen(t.completed_at) + '</td>' +
      '</tr>';
  }).join('');
  var doneTable = data.tasks.length
    ? '<table><tr><th>Court</th><th>FIPS</th><th>Type</th><th>Date range</th><th>Completed</th></tr>' + doneRows + '</table>'
    : '<div class="empty">No completed tasks yet</div>';

  var totalPages = Math.max(1, Math.ceil(data.total / data.per_page));
  var page = data.page;
  var prevDisabled = page <= 0 ? 'disabled' : '';
  var nextDisabled = (page + 1) >= totalPages ? 'disabled' : '';
  var pager = '<div class="pager">' +
    '<button onclick="changePage(-1)" ' + prevDisabled + '>&larr; Newer</button>' +
    '<span>Page ' + (page + 1) + ' of ' + totalPages + '</span>' +
    '<button onclick="changePage(1)" ' + nextDisabled + '>Older &rarr;</button>' +
    '</div>';

  var dCount = lastStatus ? lastStatus.courts.district.completed_count : 0;
  var cCount = lastStatus ? lastStatus.courts.circuit.completed_count : 0;
  document.getElementById('completed-card').innerHTML =
    '<h2>Completed tasks</h2>' +
    '<div class="stats">' +
      '<div class="stat"><div class="n">' + data.total + '</div><div class="l">Total completed</div></div>' +
      '<div class="stat"><div class="n">' + dCount + '</div><div class="l">District</div></div>' +
      '<div class="stat"><div class="n">' + cCount + '</div><div class="l">Circuit</div></div>' +
    '</div>' + doneTable + pager;
}
function loadCompleted() {
  fetch('/api/completed?page=' + completedPage)
    .then(function(r){ return r.json(); })
    .then(function(data){
      // Clamp if the last page shrank (e.g. fewer tasks than expected)
      var totalPages = Math.max(1, Math.ceil(data.total / data.per_page));
      if (completedPage > totalPages - 1) {
        completedPage = totalPages - 1;
        return loadCompleted();
      }
      renderCompleted(data);
    });
}
function changePage(delta) {
  completedPage = Math.max(0, completedPage + delta);
  loadCompleted();
}
var pendingPage = 0;
function renderPending(data) {
  var rows = data.tasks.map(function(t) {
    return '<tr>' +
      '<td>' + t.court + '</td>' +
      '<td>' + t.fips + '</td>' +
      '<td>' + t.case_type + '</td>' +
      '<td>' + (t.end_date || '') + ' → ' + (t.start_date || '') + '</td>' +
      '</tr>';
  }).join('');
  var tbl = data.tasks.length
    ? '<table><tr><th>Court</th><th>FIPS</th><th>Type</th><th>Date range</th></tr>' + rows + '</table>'
    : '<div class="empty">No scheduled tasks</div>';

  var totalPages = Math.max(1, Math.ceil(data.total / data.per_page));
  var page = data.page;
  var pager = '<div class="pager">' +
    '<button onclick="changePendingPage(-1)" ' + (page <= 0 ? 'disabled' : '') + '>&larr; Newer</button>' +
    '<span>Page ' + (page + 1) + ' of ' + totalPages + '</span>' +
    '<button onclick="changePendingPage(1)" ' + ((page + 1) >= totalPages ? 'disabled' : '') + '>Older &rarr;</button>' +
    '</div>';

  var dCount = lastStatus ? lastStatus.courts.district.pending_count : 0;
  var cCount = lastStatus ? lastStatus.courts.circuit.pending_count : 0;
  document.getElementById('pending-card').innerHTML =
    '<h2>Scheduled tasks</h2>' +
    '<div class="stats">' +
      '<div class="stat"><div class="n">' + data.total + '</div><div class="l">Total pending</div></div>' +
      '<div class="stat"><div class="n">' + dCount + '</div><div class="l">District</div></div>' +
      '<div class="stat"><div class="n">' + cCount + '</div><div class="l">Circuit</div></div>' +
    '</div>' + tbl + pager;
}
function loadPending() {
  fetch('/api/pending?page=' + pendingPage)
    .then(function(r){ return r.json(); })
    .then(function(data){
      var totalPages = Math.max(1, Math.ceil(data.total / data.per_page));
      if (pendingPage > totalPages - 1) {
        pendingPage = totalPages - 1;
        return loadPending();
      }
      renderPending(data);
    });
}
function changePendingPage(delta) {
  pendingPage = Math.max(0, pendingPage + delta);
  loadPending();
}
var ACTIVE_INTERVAL = 1000;
var IDLE_INTERVAL = 10000;
var refreshTimer = null;
function scheduleNext(ms) {
  if (refreshTimer) clearTimeout(refreshTimer);
  refreshTimer = setTimeout(refresh, ms);
}
function refresh() {
  // Cancel any pending tick so manual refreshes (e.g. from the worker control
  // or task scheduler) restart the single loop instead of spawning extra ones.
  if (refreshTimer) { clearTimeout(refreshTimer); refreshTimer = null; }
  fetch('/api/status').then(function(r){ return r.json(); }).then(function(d){
    lastStatus = d;
    document.getElementById('grid').innerHTML =
      courtCard('District', d.courts.district) + courtCard('Circuit', d.courts.circuit);
    loadCompleted();
    loadPending();
    // Poll fast while any worker is active, slowly when everything is idle
    var anyActive = d.courts.district.active.length + d.courts.circuit.active.length > 0;
    var interval = anyActive ? ACTIVE_INTERVAL : IDLE_INTERVAL;
    document.getElementById('meta').textContent =
      'Updated ' + new Date(d.generated_at).toLocaleTimeString() +
      ' · stale after ' + d.stale_threshold_seconds + 's · auto-refresh ' +
      (interval / 1000) + 's · ~ = estimate';
    scheduleNext(interval);
  }).catch(function(e){
    document.getElementById('meta').textContent = 'Error fetching status: ' + e;
    scheduleNext(IDLE_INTERVAL);
  });
}
loadCourts();
refresh();
</script>
</body>
</html>"""


DATA_PAGE = """<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<title>VA Court Scraper - Collected Data</title>
<style>
  body { font-family: -apple-system, Segoe UI, Roboto, sans-serif; margin: 24px; background: #0f1115; color: #e6e6e6; }
  h1 { font-size: 20px; margin: 0 0 4px; }
  h2 { font-size: 16px; margin: 0 0 12px; }
  .top { display: flex; align-items: center; justify-content: space-between; }
  .link { background: #21262d; color: #e6e6e6; border: 1px solid #30363d; border-radius: 6px; padding: 6px 12px; cursor: pointer; text-decoration: none; font-size: 13px; }
  .link:hover { background: #30363d; }
  .meta { color: #8a8f98; font-size: 13px; margin-bottom: 20px; }
  .grid { display: flex; gap: 24px; flex-wrap: wrap; }
  .box { flex: 1; min-width: 420px; background: #171a21; border: 1px solid #262b36; border-radius: 8px; padding: 16px; }
  table { width: 100%; border-collapse: collapse; font-size: 13px; }
  th, td { text-align: left; padding: 6px 8px; border-bottom: 1px solid #262b36; }
  th { color: #8a8f98; font-weight: 500; font-size: 11px; text-transform: uppercase; }
  .batch-item { border-bottom: 1px solid #262b36; }
  .batch-head { display: flex; align-items: center; gap: 10px; padding: 9px 4px; cursor: pointer; }
  .batch-head:hover { background: #1f242c; }
  .batch-head.open { background: #21262d; }
  .caret { color: #8a8f98; width: 12px; display: inline-block; }
  .bname { font-weight: 600; }
  .bmeta { color: #8a8f98; font-size: 12px; flex: 1; }
  .bcount { color: #8a8f98; font-size: 12px; white-space: nowrap; }
  .export { background: #21262d; color: #e6e6e6; border: 1px solid #30363d; border-radius: 6px; padding: 3px 10px; font-size: 12px; text-decoration: none; white-space: nowrap; }
  .export:hover { background: #30363d; }
  .batch-detail { padding: 4px 4px 12px 26px; }
  .cfilter { display: flex; gap: 8px; margin: 2px 0 10px; }
  .cfilter input { background: #0f1115; color: #e6e6e6; border: 1px solid #30363d; border-radius: 6px; padding: 5px 9px; font-size: 13px; min-width: 200px; }
  .cfilter input::placeholder { color: #8b949e; }
  /* Fixed layout so all columns fit without side-scrolling; long values are
     clipped with an ellipsis (full text is in each cell's tooltip). */
  .ctable { table-layout: fixed; width: 100%; }
  .ctable th, .ctable td { overflow: hidden; text-overflow: ellipsis; white-space: nowrap; }
  th.sortable { cursor: pointer; user-select: none; }
  th.sortable:hover { color: #e6e6e6; }
  .search { background: #0f1115; color: #e6e6e6; border: 1px solid #30363d; border-radius: 6px; padding: 6px 8px; font-size: 13px; margin-bottom: 12px; width: 220px; }
  .empty { color: #8a8f98; font-style: italic; padding: 8px; }
  .pager { display: flex; align-items: center; gap: 12px; margin-top: 12px; font-size: 13px; color: #8a8f98; }
  .pager button { background: #21262d; color: #e6e6e6; border: 1px solid #30363d; border-radius: 6px; padding: 5px 12px; cursor: pointer; }
  .pager button[disabled] { opacity: .4; cursor: default; }
  .muted { color: #8a8f98; }
  .ccheck { background: #1f6feb; color: #fff; border: 1px solid #388bfd; border-radius: 6px; padding: 5px 12px; font-size: 13px; cursor: pointer; margin-left: auto; }
  .ccheck:hover { background: #388bfd; }
  .ccheck[disabled] { opacity: .5; cursor: default; }
  .conflicts { margin: 2px 0 14px; }
  .conflicts h3 { font-size: 13px; margin: 0 0 8px; color: #e6e6e6; }
  .cf-note { color: #8a8f98; font-style: italic; padding: 6px 0; }
  .badge { display: inline-block; padding: 1px 7px; border-radius: 10px; font-size: 11px; font-weight: 600; white-space: nowrap; }
  .badge.exact { background: rgba(248,81,73,.18); color: #f85149; border: 1px solid rgba(248,81,73,.4); }
  .badge.review { background: rgba(210,153,34,.18); color: #d29922; border: 1px solid rgba(210,153,34,.4); }
  .badge.adv { background: rgba(248,81,73,.18); color: #f85149; border: 1px solid rgba(248,81,73,.4); margin-left: 4px; }
</style>
</head>
<body>
  <div class="top">
    <h1>Virginia Court Scraper &mdash; Collected Data</h1>
    <a class="link" href="/">&larr; Dashboard</a>
  </div>
  <div class="meta" id="meta">Loading&hellip;</div>
  <div class="box">
    <h2>Task batches</h2>
    <input type="text" id="batch-search" class="search" placeholder="Search by name&hellip;" oninput="renderBatches()">
    <div id="batches"></div>
  </div>
<script>
function fmtWhen(iso) {
  if (!iso) return '';
  var d = new Date(iso);
  var secs = Math.floor((Date.now() - d.getTime()) / 1000);
  if (secs < 60) return secs + 's ago';
  if (secs < 3600) return Math.floor(secs / 60) + 'm ago';
  if (secs < 86400) return Math.floor(secs / 3600) + 'h ago';
  return Math.round(secs / 86400) + 'd ago';
}
var batchesById = {};
var allBatches = [];
// Restore the batch that was open before a reload so its cached conflict results
// reappear when returning from the main dashboard.
var currentBatch = null;
try {
  var _ob = localStorage.getItem('cc_open_batch');
  if (_ob) { var _obn = parseInt(_ob, 10); if (!isNaN(_obn)) currentBatch = _obn; }
} catch (e) {}
var batchPage = 0;
var filterSubtype = '';
var filterJudgement = '';
function loadBatches() {
  fetch('/api/batches').then(function(r){ return r.json(); }).then(function(d){
    allBatches = d.batches || [];
    batchesById = {};
    allBatches.forEach(function(b){ batchesById[b.id] = b; });
    renderBatches();
  });
}
var lastStructSig = null;
function visibleBatches() {
  var q = (document.getElementById('batch-search').value || '').toLowerCase().trim();
  return allBatches.filter(function(b){
    return !q || (b.name && b.name.toLowerCase().indexOf(q) !== -1);
  });
}
function batchMeta(b) {
  return b.court_type + ' ' + b.case_type + ' &middot; ' +
    (b.end_date || '') + ' &rarr; ' + (b.start_date || '') + ' &middot; ' + fmtWhen(b.created_at);
}
function renderBatches() {
  var q = (document.getElementById('batch-search').value || '').toLowerCase().trim();
  var list = visibleBatches();
  // Rebuild the DOM only when the structure changes (which batches show and
  // which one is open). On plain data refreshes we update fields in place so
  // the open batch's detail (and the case table) never flashes.
  var structSig = currentBatch + '|' + list.map(function(b){ return b.id; }).join(',');
  if (structSig === lastStructSig) {
    list.forEach(function(b){
      var item = document.getElementById('bitem-' + b.id);
      if (!item) return;
      var cnt = item.querySelector('.bcount');
      if (cnt) cnt.textContent = b.cases + ' cases';
      var meta = item.querySelector('.bmeta');
      if (meta) meta.innerHTML = batchMeta(b);
    });
    document.getElementById('meta').textContent =
      list.length + ' batch(es)' + (q ? ' matching "' + q + '"' : '');
    if (currentBatch !== null) { loadCases(); restoreConflicts(currentBatch); }
    return;
  }
  lastStructSig = structSig;
  var html = list.map(function(b){
    var open = (b.id === currentBatch);
    var head = '<div class="batch-head' + (open ? ' open' : '') + '" onclick="toggleBatch(' + b.id + ')">' +
      '<span class="caret">' + (open ? '&#9662;' : '&#9656;') + '</span>' +
      '<span class="bname">' + (b.name || '&mdash;') + '</span>' +
      '<span class="bmeta">' + batchMeta(b) + '</span>' +
      '<span class="bcount">' + b.cases + ' cases</span>' +
      '<a class="export" href="#" onclick="event.stopPropagation();exportBatch(' + b.id + ',&#39;xlsx&#39;);return false;">Download Excel</a>' +
      '<a class="export" href="#" onclick="event.stopPropagation();exportBatch(' + b.id + ',&#39;csv&#39;);return false;">Download CSV</a></div>';
    var detail = '';
    if (open) {
      detail = '<div class="batch-detail" id="detail-' + b.id + '">' +
        '<div class="cfilter">' +
          '<input id="f-subtype" placeholder="Filter case type&hellip;" value="' + esc(filterSubtype) + '" onchange="applyFilter()">' +
          '<input id="f-judgement" placeholder="Filter judgement&hellip;" value="' + esc(filterJudgement) + '" onchange="applyFilter()">' +
          '<button class="ccheck" onclick="checkConflicts(' + b.id + ')">Check client conflicts</button>' +
        '</div>' +
        '<div class="conflicts" id="conflicts-body-' + b.id + '"></div>' +
        '<div id="cases-body-' + b.id + '">Loading&hellip;</div></div>';
    }
    return '<div class="batch-item" id="bitem-' + b.id + '">' + head + detail + '</div>';
  }).join('');
  document.getElementById('batches').innerHTML = html ||
    '<div class="empty">' + (allBatches.length ? 'No batches match your search' : 'No task batches yet') + '</div>';
  document.getElementById('meta').textContent = list.length + ' batch(es)' + (q ? ' matching "' + q + '"' : '');
  if (currentBatch !== null) { loadCases(); restoreConflicts(currentBatch); }
}
function toggleBatch(id) {
  if (currentBatch === id) {
    currentBatch = null;
  } else {
    currentBatch = id;
    batchPage = 0;
    filterSubtype = '';
    filterJudgement = '';
  }
  try { localStorage.setItem('cc_open_batch', currentBatch === null ? '' : String(currentBatch)); } catch (e) {}
  renderBatches();
}
var sortCol = null;
var sortDir = 'asc';
function esc(v) {
  if (v === null || v === undefined) return '';
  return String(v).replace(/&/g, '&amp;').replace(/</g, '&lt;')
                  .replace(/>/g, '&gt;').replace(/"/g, '&quot;');
}
function cell(v) {
  // Clipped to the column width by CSS; the tooltip carries the full value.
  var s = esc(v);
  return '<td title="' + s + '">' + (s || '&mdash;') + '</td>';
}
function th(col, label) {
  var arrow = (sortCol === col) ? (sortDir === 'asc' ? ' &#9650;' : ' &#9660;') : '';
  return '<th class="sortable" onclick="sortCases(&#39;' + col + '&#39;)">' + label + arrow + '</th>';
}
function sortCases(col) {
  if (sortCol === col) {
    sortDir = (sortDir === 'asc') ? 'desc' : 'asc';
  } else {
    sortCol = col;
    sortDir = 'asc';
  }
  batchPage = 0;
  loadCases();
}
function loadCases() {
  if (currentBatch === null) return;
  if (!document.getElementById('cases-body-' + currentBatch)) return;
  var url = '/api/batches/' + currentBatch + '/cases?page=' + batchPage;
  if (sortCol) url += '&sort=' + sortCol + '&dir=' + sortDir;
  if (filterSubtype) url += '&subtype=' + encodeURIComponent(filterSubtype);
  if (filterJudgement) url += '&judgement=' + encodeURIComponent(filterJudgement);
  fetch(url)
    .then(function(r){ return r.json(); }).then(function(d){
      var rows = (d.cases || []).map(function(c){
        return '<tr>' + cell(c.case_id) + cell(c.case_number) + cell(c.case_type) +
          cell(c.case_subtype) + cell(c.judgement) + cell(c.fips) + cell(c.case_date) +
          cell(c.plaintiff) + cell(c.plaintiff_attorney) + cell(c.defendant) +
          cell(c.defendant_attorney) + '</tr>';
      }).join('');
      var header = th('case_id', 'Case id') + th('case_number', 'Case number') +
        th('case_type', 'Category') + th('case_subtype', 'Case type') +
        th('judgement', 'Judgement') + th('fips', 'FIPS') + th('case_date', 'Case date') +
        th('plaintiff', 'Plaintiff') + th('plaintiff_attorney', 'Plaintiff attorney') +
        th('defendant', 'Defendant') + th('defendant_attorney', 'Defendant attorney');
      var widths = [6, 12, 7, 10, 11, 5, 8, 11, 10, 11, 9];
      var cols = '<colgroup>' + widths.map(function(w){
        return '<col style="width:' + w + '%">';
      }).join('') + '</colgroup>';
      var table = d.cases.length
        ? '<table class="ctable">' + cols + '<tr>' + header + '</tr>' + rows + '</table>'
        : '<div class="empty">No cases recorded for this batch yet</div>';
      var totalPages = Math.max(1, Math.ceil(d.total / d.per_page));
      var pager = '<div class="pager">' +
        '<button onclick="changeBatchPage(-1)" ' + (d.page <= 0 ? 'disabled' : '') + '>&larr; Prev</button>' +
        '<span>Page ' + (d.page + 1) + ' of ' + totalPages + ' (' + d.total + ' cases)</span>' +
        '<button onclick="changeBatchPage(1)" ' + ((d.page + 1) >= totalPages ? 'disabled' : '') + '>Next &rarr;</button>' +
        '</div>';
      var t = document.getElementById('cases-body-' + currentBatch);
      if (t) t.innerHTML = table + pager;
    });
}
function applyFilter() {
  var s = document.getElementById('f-subtype');
  var j = document.getElementById('f-judgement');
  filterSubtype = s ? s.value.trim() : '';
  filterJudgement = j ? j.value.trim() : '';
  batchPage = 0;
  loadCases();
}
function exportBatch(id, fmt) {
  var url = '/api/batches/' + id + '/export' + (fmt === 'xlsx' ? '.xlsx' : '');
  var qs = [];
  if (id === currentBatch) {
    if (sortCol) { qs.push('sort=' + sortCol); qs.push('dir=' + sortDir); }
    if (filterSubtype) qs.push('subtype=' + encodeURIComponent(filterSubtype));
    if (filterJudgement) qs.push('judgement=' + encodeURIComponent(filterJudgement));
  }
  window.location = url + (qs.length ? '?' + qs.join('&') : '');
}
function changeBatchPage(delta) {
  batchPage = Math.max(0, batchPage + delta);
  loadCases();
}
// Conflict results are cached in localStorage so they survive navigating to the
// main dashboard and back (each is a full page load, which wipes in-memory state).
var CONFLICT_STORE_KEY = 'cc_conflicts';
function loadConflictCache() {
  try { return JSON.parse(localStorage.getItem(CONFLICT_STORE_KEY) || '{}') || {}; }
  catch (e) { return {}; }
}
function saveConflictResult(id, data) {
  var c = loadConflictCache();
  c[id] = { data: data, checkedAt: Date.now() };
  try { localStorage.setItem(CONFLICT_STORE_KEY, JSON.stringify(c)); } catch (e) {}
}
function getConflictResult(id) {
  return loadConflictCache()[id] || null;
}
function renderConflicts(id, d, checkedAt) {
  var box = document.getElementById('conflicts-body-' + id);
  if (!box) return;
  if (!d.ok) { box.innerHTML = '<div class="cf-note">Error: ' + esc(d.error) + '</div>'; return; }
  if (d.note) { box.innerHTML = '<div class="cf-note">' + esc(d.note) + '</div>'; return; }
  var when = checkedAt ? ' &middot; checked ' + fmtWhen(new Date(checkedAt).toISOString()) : '';
  var cf = d.conflicts || [];
  if (!cf.length) {
    box.innerHTML = '<div class="cf-note">No conflicts found &mdash; scanned ' + d.cases_scanned +
      ' case(s) against ' + d.clients_checked + ' VPLC client(s).' + when + '</div>';
    return;
  }
  var rows = cf.map(function(x){
    var badge = '<span class="badge ' + x.level + '">' + x.score + '% ' + x.level + '</span>';
    var adv = x.client_adverse ? '<span class="badge adv">adverse</span>' : '';
    var clientRole = x.client_role ? ' <span class="muted">(' + esc(x.client_role) + ')</span>' : '';
    return '<tr>' +
      '<td>' + badge + '</td>' +
      cell(x.case_number) + cell(x.party_role) + cell(x.party_name) +
      '<td title="' + esc(x.client_name) + '">' + esc(x.client_name) + clientRole + adv + '</td>' +
      cell(x.elh_case_number) +
      '</tr>';
  }).join('');
  box.innerHTML = '<h3>' + cf.length + ' potential conflict(s) &mdash; ' + d.cases_scanned +
    ' cases vs ' + d.clients_checked + ' clients' + when +
    '<a class="export" style="margin-left:10px;" href="/api/batches/' + id + '/conflicts.xlsx">Download Excel</a>' +
    '<a class="export" style="margin-left:6px;" href="/api/batches/' + id + '/conflicts.csv">Download CSV</a></h3>' +
    '<table><tr><th>Match</th><th>Case number</th><th>Party role</th>' +
    '<th>Party name</th><th>VPLC client</th><th>ELH case</th></tr>' + rows + '</table>';
}
// Re-render a batch's previously computed results (e.g. after a page reload),
// but only when the panel isn't already showing this session's results.
function restoreConflicts(id) {
  if (id === null || id === undefined) return;
  var box = document.getElementById('conflicts-body-' + id);
  if (!box || box.innerHTML.trim() !== '') return;
  var stored = getConflictResult(id);
  if (stored) renderConflicts(id, stored.data, stored.checkedAt);
}
function checkConflicts(id) {
  var box = document.getElementById('conflicts-body-' + id);
  if (!box) return;
  var btn = document.querySelector('#detail-' + id + ' .ccheck');
  if (btn) { btn.disabled = true; btn.textContent = 'Checking…'; }
  box.innerHTML = '<div class="cf-note">Checking batch parties against the VPLC client list…</div>';
  function done() { if (btn) { btn.disabled = false; btn.textContent = 'Check client conflicts'; } }
  fetch('/api/batches/' + id + '/conflicts')
    .then(function(r){ return r.json(); })
    .then(function(d){
      done();
      if (d.ok) saveConflictResult(id, d);
      renderConflicts(id, d, Date.now());
    })
    .catch(function(e){ done(); box.innerHTML = '<div class="cf-note">Error: ' + e + '</div>'; });
}
loadBatches();
setInterval(loadBatches, 10000);
</script>
</body>
</html>"""


def open_browser():
    webbrowser.open('http://127.0.0.1:%d/' % PORT)


if __name__ == '__main__':
    # Open the browser shortly after the server starts. The reloader would
    # otherwise trigger this twice, so it is disabled.
    threading.Timer(1.0, open_browser).start()
    app.run(host='127.0.0.1', port=PORT, debug=False, use_reloader=False)
